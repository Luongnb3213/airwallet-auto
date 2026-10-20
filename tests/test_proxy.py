import contextlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from src.core.airwallet_proxy import parse_proxy, reserve_proxy, prepare_account_proxy
from scripts.proxy_switcher import Phone, apply_proxy
import airwallet_full_flow as full
import airwallet_dom_runner as runner
import airwallet_profile_from_excel as reader

PROXY = 'proxy.example:8080:synthetic-user:synthetic-secret'
PROFILE = {'phone_number': '07012345678', 'account_password': 'TestOnly123', 'kana_sei': 'テスト',
           'kana_mei': 'タロウ', 'birthday': '1999-02-03', 'pin': '0042'}


class ProxyAssignmentTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'accounts.xlsx'

    def fixture(self, count=2):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Inputs'
        sheet.append(['phone_number', 'pin', 'status'])
        for phone in ('07012345678', '07012345679', '07012345670'):
            sheet.append([phone, '0042', 'PENDING'])
        pool = book.create_sheet('Proxies')
        pool.append(['proxy', 'status', 'proxy_id'])
        for index in range(count):
            pool.append([f'proxy.example:8080:synthetic-user-{index}:synthetic-secret', 'ACTIVE', ''])
        book.save(self.path)
        book.close()

    def test_each_account_gets_one_distinct_proxy_and_resume_keeps_it(self):
        self.fixture()
        first = reserve_proxy(self.path, 'Inputs', row=2)
        second = reserve_proxy(self.path, 'Inputs', row=3)
        resumed = reserve_proxy(self.path, 'Inputs', phone_number='+817012345678')
        self.assertNotEqual(first['proxy_id'], second['proxy_id'])
        self.assertEqual(first['proxy_id'], resumed['proxy_id'])
        book = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual(book['Inputs']['B2'].value, '0042')
            self.assertEqual(book['Proxies']['D2'].value, '07012345678')
            self.assertNotIn('synthetic-secret', book['Inputs']['E2'].value)
        finally:
            book.close()

    def test_exhausted_pool_never_wraps_or_reuses_another_account_proxy(self):
        self.fixture(count=1)
        reserve_proxy(self.path, 'Inputs', row=2)
        with self.assertRaisesRegex(ValueError, 'Hết proxy'):
            reserve_proxy(self.path, 'Inputs', row=3)

    def test_manual_proxy_change_after_assignment_is_rejected(self):
        self.fixture()
        first = reserve_proxy(self.path, 'Inputs', row=2)
        book = openpyxl.load_workbook(self.path)
        book['Inputs']['D2'] = 'PX-000002'
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'không khớp'):
            reserve_proxy(self.path, 'Inputs', row=2)

    def test_duplicate_raw_proxies_are_rejected_without_credentials_in_error(self):
        self.fixture()
        book = openpyxl.load_workbook(self.path)
        book['Proxies']['A3'] = book['Proxies']['A2'].value
        book.save(self.path)
        book.close()
        with self.assertRaises(ValueError) as caught:
            reserve_proxy(self.path, 'Inputs', row=2)
        self.assertNotIn('synthetic-secret', str(caught.exception))

    def test_duplicate_phone_rows_are_rejected_before_assignment(self):
        self.fixture()
        book = openpyxl.load_workbook(self.path)
        book['Inputs']['A3'] = book['Inputs']['A2'].value
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'đúng một dòng'):
            reserve_proxy(self.path, 'Inputs', row=2)

    def test_wrong_row_and_phone_cannot_claim_proxy(self):
        self.fixture()
        with self.assertRaisesRegex(ValueError, 'đúng một dòng'):
            reserve_proxy(self.path, 'Inputs', row=2, phone_number='07012345679')

    def test_inactive_proxy_is_not_used(self):
        self.fixture(count=1)
        book = openpyxl.load_workbook(self.path)
        book['Proxies']['B2'] = 'INACTIVE'
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'Hết proxy'):
            reserve_proxy(self.path, 'Inputs', row=2)

    def test_device_failure_keeps_same_assignment_and_records_sanitized_error(self):
        self.fixture()
        with patch('scripts.proxy_switcher.Phone', side_effect=RuntimeError('synthetic-secret')):
            with self.assertRaisesRegex(RuntimeError, 'Không áp dụng'):
                prepare_account_proxy(self.path, 'Inputs', 'test-device', row=2)
        assigned = reserve_proxy(self.path, 'Inputs', row=2)
        self.assertEqual(assigned['proxy_id'], 'PX-000001')
        book = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual(book['Proxies']['B2'].value, 'ASSIGNED')
            self.assertNotIn('synthetic-secret', book['Proxies']['G2'].value)
        finally:
            book.close()

    def test_success_records_ip_without_exposing_credentials_in_result(self):
        self.fixture()
        with patch('scripts.proxy_switcher.Phone'), \
             patch('scripts.proxy_switcher.apply_proxy', return_value={'ok': True, 'ip': '192.0.2.1'}):
            result = prepare_account_proxy(self.path, 'Inputs', 'test-device', row=2)
        self.assertEqual(result['proxy_id'], 'PX-000001')
        self.assertNotIn('synthetic-secret', json.dumps(result))
        book = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual(book['Proxies']['E2'].value, '192.0.2.1')
        finally:
            book.close()

    def test_parallel_processes_cannot_assign_the_same_proxy_twice(self):
        self.fixture(count=1)
        code = '''import sys,json
from src.core.airwallet_proxy import reserve_proxy
try:
    result=reserve_proxy(sys.argv[1], 'Inputs', row=int(sys.argv[2]))
    print(json.dumps({'ok':True,'proxy_id':result['proxy_id']}))
except ValueError:
    print(json.dumps({'ok':False}))
'''
        jobs = [subprocess.Popen([sys.executable, '-c', code, str(self.path), str(row)], cwd=ROOT,
                                 stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True) for row in (2, 3)]
        results = []
        for job in jobs:
            output, error = job.communicate(timeout=25)
            self.assertEqual(job.returncode, 0, error)
            results.append(json.loads(output))
        self.assertEqual(sum(result['ok'] for result in results), 1)


class ProxyDeviceTests(unittest.TestCase):
    def test_parser_preserves_colons_in_password_without_leaking_bad_input(self):
        self.assertEqual(parse_proxy(PROXY + ':suffix')['password'], 'synthetic-secret:suffix')
        for raw in ('secret', 'host:0:user:synthetic-secret', 'host:65536:user:synthetic-secret', 'host:8080::synthetic-secret'):
            with self.assertRaises(ValueError) as caught:
                parse_proxy(raw)
            self.assertNotIn('synthetic-secret', str(caught.exception))

    def test_apply_uses_only_selected_proxy_and_accepts_same_exit_ip(self):
        phone = Mock()
        phone.helper.return_value = True
        phone.exit_ip.return_value = '192.0.2.1'
        selected = parse_proxy(PROXY)
        result = apply_proxy(phone, selected)
        self.assertTrue(result['ok'])
        phone.set_proxy.assert_called_once_with(selected)
        phone.enable.assert_called_once_with(selected, True)
        phone.exit_ip.assert_called_once_with()

    def test_unready_proxy_does_not_switch_to_a_different_proxy(self):
        phone = Mock()
        phone.helper.return_value = False
        with patch('scripts.proxy_switcher.time.monotonic', side_effect=[0, 0, 31]), \
             patch('scripts.proxy_switcher.time.sleep'):
            result = apply_proxy(phone, parse_proxy(PROXY))
        self.assertFalse(result['ok'])
        phone.set_proxy.assert_called_once()
        phone.exit_ip.assert_not_called()

    def test_auto_device_selection_requires_exactly_one_usb_phone(self):
        devices = [SimpleNamespace(type='usb', id='first'), SimpleNamespace(type='usb', id='second')]
        with patch('frida.enumerate_devices', return_value=devices):
            with self.assertRaisesRegex(RuntimeError, 'đúng một iPhone'):
                Phone('auto')

    def test_proxy_failure_stops_full_flow_before_crane_and_signup(self):
        with patch.object(sys, 'argv', ['full', '--xlsx', 'unused.xlsx', '--row', '2']), \
             patch.object(reader, 'load_row', return_value=(PROFILE, 2, 'Inputs')), \
             patch('flow.build_otp_client', return_value=object()), \
             patch.object(full, 'prepare_account_proxy', side_effect=RuntimeError('proxy not ready')), \
             patch.object(full, 'run_cmd_output') as crane, \
             patch.object(runner, 'run_with_frida') as signup, \
             contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(full.main(), 1)
        crane.assert_not_called()
        signup.assert_not_called()


if __name__ == '__main__':
    unittest.main()
