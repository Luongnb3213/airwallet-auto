from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from airwallet_profile_from_excel import load_row, profile_from_record
from airwallet_full_flow import write_row_status
from src.connections.xlsx_connection import INPUT_HEADERS, PROXIES_HEADERS, XlsxConnection
from crane_container_manager import CRANE_COLUMNS, write_crane_assignment
from airwallet_batch_flow import runnable_rows


class LineCoinDataTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'inputs.xlsx'

    def fixture(self, headers, values):
        book = openpyxl.Workbook()
        book.active.title = 'Inputs'
        book.active.append(headers)
        book.active.append(values)
        book.save(self.path)
        book.close()

    def test_two_password_columns_remain_distinct_in_either_order(self):
        for keys, values in [
            (['email', 'password', 'name', 'account_password'], ['test@example.com', 'first@123', 'Example', 'second!789']),
            (['account_password', 'name', 'password', 'email'], ['second!789', 'Example', 'first@123', 'test@example.com']),
        ]:
            with self.subTest(headers=keys):
                self.fixture(keys, values)
                record, row, sheet = load_row(self.path, 'Inputs', 2)
                profile = profile_from_record(record)
                self.assertEqual(profile['password'], 'first@123')
                self.assertEqual(profile['account_password'], 'second!789')
                self.assertEqual(profile['name'], 'Example')

    def test_status_writer_uses_line_coin_message_column(self):
        headers = ['email', 'password', 'name', 'account_password', 'status', 'message']
        self.fixture(headers, ['test@example.com', 'one', 'Example', 'two', 'PENDING', 'old'])
        write_row_status(self.path, 'Inputs', 2, 'FAILED', 'test error')
        book = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual([c.value for c in book.active[1]], headers)
            self.assertEqual(book.active['E2'].value, 'FAILED')
            self.assertEqual(book.active['F2'].value, 'test error')
            self.assertEqual(book.active['B2'].value, 'one')
            self.assertEqual(book.active['D2'].value, 'two')
        finally:
            book.close()

    def test_legacy_error_details_workbooks_still_write_existing_column(self):
        self.fixture(['email', 'status', 'error_details'], ['test@example.com', 'PENDING', 'old'])
        write_row_status(self.path, 'Inputs', 2, 'SUCCESS', '')
        book = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual(book.active.max_column, 3)
            self.assertEqual(book.active['B2'].value, 'SUCCESS')
            self.assertIsNone(book.active['C2'].value)
        finally:
            book.close()

    def test_workbook_has_schema_and_text_account_identifiers(self):
        self.assertTrue(XlsxConnection.create_template(str(self.path)))
        book = openpyxl.load_workbook(self.path, read_only=True)
        try:
            self.assertEqual(book.sheetnames, ['Inputs', 'Proxies'])
            self.assertEqual([c.value for c in next(book.active.rows)], INPUT_HEADERS)
            for row in book['Inputs'].iter_rows(min_row=2):
                for column in (4, 8):  # phone_number and PIN remain text in operational data.
                    if row[column].value is not None:
                        self.assertIsInstance(row[column].value, str)
            self.assertEqual([c.value for c in next(book['Proxies'].rows)], PROXIES_HEADERS)
            for sheet in book:
                self.assertTrue(all(cell.value is None for row in sheet.iter_rows(min_row=2) for cell in row))
        finally:
            book.close()

    def test_template_creation_preserves_an_existing_workbook(self):
        self.assertEqual(INPUT_HEADERS[:4], ['email', 'password', 'name', 'account_password'])
        self.assertEqual(INPUT_HEADERS[4:9], ['phone_number', 'kana_sei', 'kana_mei', 'birthday', 'pin'])
        self.fixture(['phone_number', 'status'], ['07012345678', 'SUCCESS'])
        original = self.path.read_bytes()
        self.assertTrue(XlsxConnection.create_template(str(self.path)))
        self.assertEqual(self.path.read_bytes(), original)

    def test_phone_only_row_without_email_is_usable(self):
        self.fixture(['phone_number', 'account_password', 'kana_sei', 'kana_mei', 'birthday', 'pin'],
                     ['07012345678', 'Example123', 'テスト', 'タロウ', '1999/02/03', '0042'])
        record, row, sheet = load_row(self.path, 'Inputs', 2)
        profile = profile_from_record(record)
        self.assertEqual(profile['phone_number'], '07012345678')
        self.assertEqual(profile['pin'], '0042')
        self.assertEqual(profile['birthday'], '1999/02/03')

    def test_template_contains_all_device_and_container_columns(self):
        self.assertTrue(set(CRANE_COLUMNS).issubset(INPUT_HEADERS))
        self.assertEqual(len(INPUT_HEADERS), len(set(INPUT_HEADERS)))

    def test_crane_write_preserves_phone_only_data_and_routes_batch(self):
        self.fixture(['phone_number', 'pin', 'status'], ['07012345678', '0042', 'PENDING'])
        for _ in range(2):
            write_crane_assignment(self.path, 'Inputs', 2, 'test-container-id', 'Test container', 'READY',
                                   'test-device-id', 'Test iPhone')
        record, row, sheet = load_row(self.path, 'Inputs', 2)
        self.assertEqual(record['phone_number'], '07012345678')
        self.assertEqual(record['pin'], '0042')
        self.assertEqual(record['crane_container_id'], 'test-container-id')
        self.assertEqual(record['frida_device_id'], 'test-device-id')
        self.assertEqual(profile_from_record(record)['frida_device_id'], 'test-device-id')
        self.assertEqual(runnable_rows(self.path, 'Inputs', 0)[1], [{'row': 2, 'device_id': 'test-device-id'}])
        book = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual(book.active.max_column, 3 + len(CRANE_COLUMNS))
            self.assertEqual([cell.value for cell in book.active[1]][:3], ['phone_number', 'pin', 'status'])
        finally:
            book.close()

    def test_device_id_alias_survives_profile_loading(self):
        self.fixture(['phone_number', 'device_id'], ['07012345678', 'test-device-id'])
        record, _, _ = load_row(self.path, 'Inputs', 2)
        self.assertEqual(profile_from_record(record)['frida_device_id'], 'test-device-id')


if __name__ == '__main__':
    unittest.main()
