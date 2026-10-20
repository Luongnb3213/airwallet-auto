"""The container-assignment race, and the lock that closes it.

Container selection is a read-decide-write across the workbook:
`assigned_container_rows()` reads which containers other rows hold,
`ensure_row()` picks one that looks free, and only the final
`write_crane_assignment()` takes `excel_write_lock`. With one `ensure-row`
process per device, two can both read "container X is free" and both claim it,
putting two accounts in one container.

These tests use threads rather than processes: the lock is a file lock, so
mutual exclusion between threads in one process exercises the same primitive
without needing a device or a second interpreter.
"""
import sys
import threading
import time
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts'))

from crane_container_manager import (
    assigned_container_rows,
    container_assignment_lock,
    write_crane_assignment,
)

HEADERS = ['email', 'phone_number', 'status', 'crane_container_id', 'crane_container_name',
           'crane_status', 'crane_assigned_at', 'crane_last_used_at',
           'frida_device_id', 'frida_device_name']


def make_workbook(path: Path, rows: int = 3) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Inputs'
    sheet.append(HEADERS)
    for index in range(rows):
        sheet.append([f'row{index}@example.test', '07000000000', 'PENDING', '', '', '', '', '', '', ''])
    book.save(path)
    book.close()
    return path


class AssignmentLockTests(unittest.TestCase):
    def setUp(self):
        import tempfile
        self._tmp = tempfile.TemporaryDirectory()
        self.xlsx = make_workbook(Path(self._tmp.name) / 'accounts.xlsx')

    def tearDown(self):
        self._tmp.cleanup()

    def test_lock_is_mutually_exclusive(self):
        order: list[str] = []
        inside = threading.Event()
        release = threading.Event()

        def first():
            with container_assignment_lock(self.xlsx):
                order.append('first-in')
                inside.set()
                release.wait(5)
                order.append('first-out')

        def second():
            inside.wait(5)
            with container_assignment_lock(self.xlsx):
                order.append('second-in')

        thread_a = threading.Thread(target=first)
        thread_b = threading.Thread(target=second)
        thread_a.start()
        thread_b.start()
        inside.wait(5)
        time.sleep(0.2)                      # second is now blocked on the lock
        self.assertEqual(order, ['first-in'], 'the second holder must not be inside yet')
        release.set()
        thread_a.join(5)
        thread_b.join(5)
        self.assertEqual(order, ['first-in', 'first-out', 'second-in'])

    def test_lock_is_released_when_the_body_raises(self):
        with self.assertRaises(RuntimeError):
            with container_assignment_lock(self.xlsx):
                raise RuntimeError('boom')
        # Must be reacquirable; a leaked lock would hang the next row.
        acquired = []
        worker = threading.Thread(target=lambda: (
            container_assignment_lock(self.xlsx).__enter__(), acquired.append(True)))
        worker.start()
        worker.join(5)
        self.assertEqual(acquired, [True])

    def test_serialised_selection_does_not_hand_one_container_to_two_rows(self):
        """With the lock held across read-decide-write, the second picks a different one."""
        pool = ['c1', 'c2']
        picked: list[str] = []
        errors: list[BaseException] = []
        barrier = threading.Barrier(2, timeout=5)

        def claim(row_number: int) -> None:
            try:
                barrier.wait()               # both threads race for the lock together
                with container_assignment_lock(self.xlsx):
                    taken = assigned_container_rows(self.xlsx, 'Inputs', row_number)
                    free = [item for item in pool if not taken.get(item)]
                    if not free:
                        picked.append('')
                        return
                    choice = free[0]
                    time.sleep(0.05)         # widen the old TOCTOU window
                    write_crane_assignment(self.xlsx, 'Inputs', row_number, choice,
                                           f'Profile {choice}', 'ASSIGNED', 'dev', 'Device')
                    picked.append(choice)
            except BaseException as error:   # surface thread failures in the test
                errors.append(error)

        threads = [threading.Thread(target=claim, args=(row,)) for row in (2, 3)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(10)

        self.assertEqual(errors, [])
        self.assertEqual(len(picked), 2)
        self.assertEqual(sorted(picked), ['c1', 'c2'],
                         'each row must get its own container')
        self.assertEqual(len(set(picked)), 2, 'no container assigned twice')

        book = openpyxl.load_workbook(self.xlsx)
        sheet = book['Inputs']
        column = HEADERS.index('crane_container_id') + 1
        stored = [sheet.cell(row=row, column=column).value for row in (2, 3)]
        book.close()
        self.assertEqual(sorted(str(value) for value in stored), ['c1', 'c2'])

    def test_unserialised_selection_can_double_assign(self):
        """Demonstrates the bug the lock fixes: same logic, no lock."""
        pool = ['c1', 'c2']
        picked: list[str] = []
        gate = threading.Barrier(2, timeout=5)

        def claim(row_number: int) -> None:
            gate.wait()
            taken = assigned_container_rows(self.xlsx, 'Inputs', row_number)
            free = [item for item in pool if not taken.get(item)]
            choice = free[0] if free else ''
            time.sleep(0.05)                 # the window both threads pass through
            if choice:
                write_crane_assignment(self.xlsx, 'Inputs', row_number, choice,
                                       f'Profile {choice}', 'ASSIGNED', 'dev', 'Device')
            picked.append(choice)

        threads = [threading.Thread(target=claim, args=(row,)) for row in (2, 3)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(10)

        self.assertEqual(len(picked), 2)
        self.assertEqual(len(set(picked)), 1,
                         'without the lock both rows read the same free container')
        self.assertEqual(set(picked), {'c1'})

    def test_ensure_row_takes_the_lock(self):
        """Guard the wiring: ensure_row must not bypass the lock."""
        import crane_container_manager as manager

        source = Path(manager.__file__).read_text(encoding='utf-8')
        self.assertIn('with container_assignment_lock(Path(args.xlsx).expanduser()):', source)
        self.assertIn('def _ensure_row_locked(', source)

    def test_assignment_lock_uses_its_own_file(self):
        """It must not reuse excel_write_lock, which is not reentrant."""
        with container_assignment_lock(self.xlsx):
            from src.connections.xlsx_connection import excel_write_lock
            # Nesting must work: write_crane_assignment does exactly this.
            with excel_write_lock(self.xlsx):
                pass
        names = {path.name for path in self.xlsx.parent.iterdir()}
        self.assertIn(f'.{self.xlsx.name}.crane-assign.lock', names)
        self.assertIn(f'.{self.xlsx.name}.write.lock', names)


if __name__ == '__main__':
    unittest.main()
