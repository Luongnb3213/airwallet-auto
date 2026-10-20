from __future__ import annotations

import argparse
import json
import os
import re
import shlex
import signal
import subprocess
import sys
import threading
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.connections.xlsx_connection import excel_write_lock, replace_excel_file  # noqa: E402
from src.core.airwallet_proxy import prepare_account_proxy  # noqa: E402


DOM_IDLE_TIMEOUT_SEC = float(os.environ.get("AIRWALLET_DOM_IDLE_TIMEOUT_SEC", "90") or 90)


def quote_cmd(cmd: list[str]) -> str:
    return " ".join(shlex.quote(str(part)) for part in cmd)


def force_kill_app(device_id: str, bundle_id: str = "") -> None:
    cmd = [sys.executable, "scripts/airwallet_dom_runner.py", "--action", "kill", "--device-id", device_id]
    if bundle_id:
        cmd.extend(["--bundle-id", bundle_id])
    try:
        subprocess.run(cmd, cwd=str(ROOT_DIR), text=True, capture_output=True, timeout=30, check=False)
        print(f"[flow] Đã force-kill app AirWallet trên máy {device_id} sau khi treo.", flush=True)
    except Exception as exc:
        print(f"[flow] Force-kill app trên {device_id} lỗi (bỏ qua): {exc}", flush=True)


def run_cmd_output(
    cmd: list[str],
    title: str,
    stream: bool = False,
    idle_timeout: float = 0.0,
    kill_app_device: str = "",
) -> str:
    print(f"\n--- {title} ---", flush=True)
    print("$ " + quote_cmd(cmd), flush=True)
    proc = subprocess.Popen(
        cmd,
        cwd=str(ROOT_DIR),
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
        start_new_session=True,
    )
    assert proc.stdout is not None
    lines: list[str] = []
    idle_fired = {"v": False}
    last_activity = [time.monotonic()]
    stop_watch = threading.Event()

    def _watch(p=proc, flag=idle_fired, last=last_activity, stop=stop_watch, limit=idle_timeout):
        while not stop.wait(5.0):
            if time.monotonic() - last[0] > limit:
                flag["v"] = True
                try:
                    os.killpg(os.getpgid(p.pid), signal.SIGKILL)
                except Exception:
                    try:
                        p.kill()
                    except Exception:
                        pass
                return

    if idle_timeout and idle_timeout > 0:
        threading.Thread(target=_watch, daemon=True).start()
    try:
        for line in proc.stdout:
            last_activity[0] = time.monotonic()
            lines.append(line)
            if stream:
                print(line, end="", flush=True)
        code = proc.wait()
    finally:
        stop_watch.set()
    output = "".join(lines)
    if idle_fired["v"]:
        print(
            f"[flow] {title}: im lặng > {idle_timeout:.0f}s -> coi như treo, đã kill cây tiến trình.",
            flush=True,
        )
        if kill_app_device:
            force_kill_app(kill_app_device)
        raise RuntimeError(f"{title} treo (im lặng > {idle_timeout:.0f}s).")
    if code != 0:
        if not stream and output.strip():
            print(output.strip(), flush=True)
        raise RuntimeError(f"{title} lỗi exit={code}.")
    return output


def parse_json_from_output(output: str) -> Any:
    decoder = json.JSONDecoder()
    candidates: list[tuple[int, Any]] = []
    for index, char in enumerate(output):
        if char != "{":
            continue
        try:
            value, end = decoder.raw_decode(output[index:])
        except json.JSONDecodeError:
            continue
        candidates.append((index + end, value))
    if not candidates:
        raise RuntimeError("Không đọc được JSON output.")
    return max(candidates, key=lambda item: item[0])[1]


def ensure_status_headers(ws) -> dict[str, int]:
    headers = [str(cell.value or "").strip().lower() for cell in ws[1]]
    if not any(headers):
        headers = ["email", "password", "name", "account_password", "status", "message"]
        for index, header in enumerate(headers, start=1):
            ws.cell(row=1, column=index, value=header)
    message_header = "message" if "message" in headers else ("error_details" if "error_details" in headers else "message")
    for header in ("status", message_header):
        if header not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=header)
            headers.append(header)
    columns = {header: index + 1 for index, header in enumerate(headers) if header}
    columns["error_details"] = columns[message_header]
    return columns


def write_row_status(xlsx: Path, sheet_name: str, row: int, status: str, error_details: str = "") -> None:
    with excel_write_lock(xlsx):
        lock_path = xlsx.parent / f".~lock.{xlsx.name}#"
        if lock_path.exists():
            print(f"[excel] Không ghi được status vì Excel đang mở/lock: {lock_path}", flush=True)
            return
        wb = openpyxl.load_workbook(xlsx)
        try:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            elif "Inputs" in wb.sheetnames:
                ws = wb["Inputs"]
            else:
                ws = wb[wb.sheetnames[0]]
            col = ensure_status_headers(ws)
            ws.cell(row=row, column=col["status"], value=status)
            ws.cell(row=row, column=col["error_details"], value=error_details)
            tmp = xlsx.with_name(f"{xlsx.stem}.{os.getpid()}.tmp.xlsx")
            wb.save(tmp)
            replace_excel_file(tmp, xlsx)
        finally:
            wb.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run one AirWallet row through SMS OTPBase registration.")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sheet", default="Inputs")
    parser.add_argument("--row", type=int, required=True)
    parser.add_argument("--no-reload", action="store_true")
    parser.add_argument("--no-container", action="store_true", help="Use the current app container without creating/switching Crane containers.")
    parser.add_argument("--reload-crane", action="store_true")
    parser.add_argument("--no-submit", action="store_true")
    parser.add_argument("--submit-final", action="store_true", help="Compatibility option; creation is enabled unless --no-submit.")
    parser.add_argument("--device-id", default=os.environ.get("FRIDA_DEVICE_ID", "auto"))
    parser.add_argument(
        "--container-mode",
        choices=["active-then-create", "active-then-next", "next-active", "create"],
        default="create",
    )
    parser.add_argument("--profile-json", default="")
    parser.add_argument("--wait-timeout-ms", type=int, default=15000)
    parser.add_argument("--max-steps", type=int, default=40)
    parser.add_argument("--dom-idle-timeout-sec", type=float, default=DOM_IDLE_TIMEOUT_SEC)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    from src.utils.secret_files import reap_stale
    # Older revisions wrote the profile (with the PIN) here and relied on a
    # `finally` that a SIGKILL'd row never reached. Reclaim any leftovers.
    stale = reap_stale(ROOT_DIR / 'agents' / 'runtime', 'input_*.json', 3600)
    if stale:
        print(f'[flow] Đã xoá {len(stale)} file profile tạm sót lại từ lần chạy trước '
              f'(chứa PIN plaintext).', flush=True)
    from flow import ensure_configured, validate_profile, load_settings, build_otp_client
    from airwallet_profile_from_excel import load_row, profile_from_record
    from airwallet_dom_runner import run_with_frida, build_parser as dom_parser
    ensure_configured()
    xlsx = Path(args.xlsx).expanduser()
    settings = load_settings()
    result = None
    try:
        record, row_number, sheet = load_row(xlsx, args.sheet, args.row)
        profile = dict(settings.get('registration_defaults') or {})
        profile.update(profile_from_record(record))
        if args.profile_json:
            raw = json.loads(Path(args.profile_json).read_text(encoding='utf-8-sig'))
            profile.update(raw.get('profile', raw))
        profile = validate_profile(profile)
        from inputs import select_device_id
        args.device_id = select_device_id(args.device_id, profile=profile, settings=settings)
        if build_otp_client(settings) is None:
            raise ValueError('Chưa cấu hình AIRWALLET_OTPBASE_API_KEY/otpbase_api_key; chưa đổi container hay gửi SMS.')
        proxy_result = prepare_account_proxy(xlsx, sheet, args.device_id, row=row_number, phone_number=profile['phone_number'])
        profile.update({key: proxy_result[key] for key in ('proxy_id', 'proxy_used')})
        print('[proxy] ' + proxy_result['proxy_id'] + ' ready, exit IP ' + proxy_result['ip'], flush=True)
        if not args.no_container:
            crane_cmd = [sys.executable, 'scripts/crane_container_manager.py', 'ensure-row',
                         '--xlsx', str(xlsx), '--sheet', sheet, '--row', str(row_number),
                         '--device-id', args.device_id, '--container-mode', args.container_mode]
            if not args.reload_crane:
                crane_cmd.append('--no-reload')
            crane_result = parse_json_from_output(run_cmd_output(crane_cmd, 'Chuẩn bị container'))
            if not isinstance(crane_result, dict) or crane_result.get('ok') is not True:
                raise RuntimeError('Crane chưa xác nhận container sẵn sàng; chưa mở app để đăng ký.')
        safe_device = re.sub(r'[^A-Za-z0-9_.-]+', '_', args.device_id)[:48]
        dom_args = ['--action', 'run', '--device-id', args.device_id,
                    '--max-steps', str(args.max_steps), '--wait-timeout-ms', str(args.wait_timeout_ms)]
        if not args.no_container:
            dom_args.append('--fresh-launch')
        if args.no_submit:
            dom_args.append('--no-submit')
        parsed_dom_args = dom_parser().parse_args(dom_args)
        parsed_dom_args.proxy_prepared = True
        # run_with_frida runs in this process, so hand the profile over in
        # memory. It used to be written to agents/runtime/input_*.json and
        # deleted in `finally` -- but airwallet_batch_flow kills a hung row with
        # SIGKILL, which no `finally` survives, leaving a plaintext PIN on disk.
        parsed_dom_args.profile_data = profile
        result = run_with_frida(parsed_dom_args)
        result['proxy'] = proxy_result
        destination = ROOT_DIR / 'dumps' / f'row-{row_number}-{safe_device}-{os.getpid()}.json'
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        if result.get('complete'):
            status = 'SUCCESS'
        elif result.get('ready_to_submit'):
            status = 'READY_TO_SUBMIT'
        else:
            status = 'FAILED'
        message = '' if status != 'FAILED' else str(result.get('message') or result.get('status'))
        write_row_status(xlsx, sheet, row_number, status, message)
        print(json.dumps({'ok': status != 'FAILED', 'status': status, 'flow_status': result.get('status'),
                          'row': row_number, 'dump': str(destination)}, ensure_ascii=False), flush=True)
        return 0 if status != 'FAILED' else 1
    except Exception as error:
        message = str(error)
        if 'profile' in locals():
            for key in ('password', 'account_password', 'pin', 'otp'):
                value = str(profile.get(key) or '')
                if value:
                    message = message.replace(value, '[redacted]')
        if xlsx.exists():
            write_row_status(xlsx, args.sheet, args.row, 'FAILED', message)
        print('[airwallet] ' + message, file=sys.stderr, flush=True)
        return 1



if __name__ == "__main__":
    raise SystemExit(main())
