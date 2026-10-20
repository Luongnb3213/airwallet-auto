from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.connections.xlsx_connection import COLUMN_ALIASES, normalize_status  # noqa: E402


PROFILE_KEYS = [
    "email",
    "name",
    "account_password",
    "phone_number",
    "kana_sei",
    "kana_mei",
    "birthday",
    "pin",
    "otp_after_seq",
    "phone_fingerprint",
    "nickname",
    "password",
    "password_confirm",
    "secret_question",
    "secret_answer",
    "dob",
    "prefecture",
    "gender",
    "referral_code",
    "otp_inbox",
    "otp_password",
    "otp_imap_host",
    "crane_container_id",
    "crane_container_name",
    "crane_status",
    "frida_device_id",
    "frida_device_name",
    "proxy_id",
    "proxy_used",
]

EXTRA_ALIASES = {
    "airwallet_email": "email",
    "mail_address": "email",
    "birth date": "birthday",
    "birth_date": "birthday",
    "birthday": "birthday",
    "phone": "phone_number",
    "phone number": "phone_number",
    "tel": "phone_number",
    "kana surname": "kana_sei",
    "kana given name": "kana_mei",
    "resident": "prefecture",
    "residence": "prefecture",
    "pref": "prefecture",
}


def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    return EXTRA_ALIASES.get(text) or COLUMN_ALIASES.get(text, text)


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_digits(value: str) -> str:
    return re.sub(r"\D", "", value or "")


def normalize_dob(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    compact = normalize_digits(raw)
    if len(compact) == 8:
        return compact

    parts = re.match(r"^\s*(\d{1,4})[\/\-.](\d{1,2})[\/\-.](\d{1,4})\s*$", raw)
    if not parts:
        return compact

    first, second, third = parts.groups()
    if len(first) == 4:
        year, month, day = first, second, third
    elif len(third) == 4:
        month, day, year = first, second, third
    else:
        return compact
    return f"{int(year):04d}{int(month):02d}{int(day):02d}"


def read_config_xlsx() -> str:
    config_path = ROOT_DIR / "config.json"
    if not config_path.exists():
        return ""
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except Exception:
        return ""
    return str(data.get("xlsx_path") or "").strip()


def read_config_sheet() -> str:
    config_path = ROOT_DIR / "config.json"
    if not config_path.exists():
        return "Inputs"
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except Exception:
        return "Inputs"
    return str(data.get("active_sheet") or "Inputs").strip() or "Inputs"


def choose_sheet(wb, sheet_name: str):
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    for fallback in ("Inputs", "Iclouds", "Gmails", "Outlooks"):
        if fallback in wb.sheetnames:
            return wb[fallback]
    return wb[wb.sheetnames[0]]


def load_row(xlsx_path: Path, sheet_name: str, row_number: int | None, email: str = "") -> tuple[dict, int, str]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        ws = choose_sheet(wb, sheet_name)
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise RuntimeError(f"Sheet {ws.title!r} is empty.")

        headers = [normalize_header(cell) for cell in header_row]
        wanted_email = email.strip().lower()

        for index, row in enumerate(rows, start=2):
            if row_number is not None and index != row_number:
                continue
            record = {}
            for pos, header in enumerate(headers):
                if not header:
                    continue
                value = cell_text(row[pos]) if pos < len(row) else ""
                if header in record and not value:
                    continue
                record[header] = value
            row_email = str(record.get("email") or "").strip()
            if not row_email and not str(record.get("phone_number") or "").strip():
                continue
            if wanted_email and row_email.lower() != wanted_email:
                continue
            if row_number is None and not wanted_email:
                status = normalize_status(record.get("status", "PENDING"))
                if status not in ("PENDING", "FAILED"):
                    continue
            return record, index, ws.title
    finally:
        wb.close()

    if row_number is not None:
        raise RuntimeError(f"No usable data found at row {row_number}.")
    if email:
        raise RuntimeError(f"No usable row found for email {email!r}.")
    raise RuntimeError("Không có dòng PENDING/FAILED chứa email hoặc phone_number.")


def profile_from_record(record: dict) -> dict:
    # Preserve values; flow.validate_profile enforces the observed iOS field rules.
    profile = {}
    for key in PROFILE_KEYS:
        value = str(record.get(key) or "").strip()
        if value:
            profile[key] = value
    return profile


def write_profile_json(profile: dict, out_path: Path) -> None:
    """Write the profile owner-only, and say out loud that it holds a PIN.

    This file is persistent, not temporary: it is the default --profile-json of
    airwallet_dom_runner.py, so a plaintext PIN and account password stay on
    disk until someone deletes them.
    """
    from src.utils.secret_files import is_owner_only, write_private_json

    written = write_private_json(out_path, profile, indent=2)
    sensitive = sorted(key for key in ("pin", "account_password", "password") if profile.get(key))
    if sensitive:
        if is_owner_only(written):
            scope = "quyền đã giới hạn về chỉ chủ sở hữu"
        else:
            scope = "CHƯA giới hạn được quyền - file đang đọc được bởi account khác"
        print(
            f"[profile] {written} chứa {', '.join(sensitive)} dạng plaintext; {scope}. "
            "Xoá sau khi dùng.",
            file=sys.stderr,
            flush=True,
        )


def main() -> int:
    parser = argparse.ArgumentParser(description="Export one Excel row into a AirWallet runtime profile JSON file.")
    parser.add_argument("--xlsx", default=read_config_xlsx(), help="Path to XLSX. Defaults to config.json xlsx_path.")
    parser.add_argument("--sheet", default=read_config_sheet(), help="Sheet name. Defaults to config active_sheet.")
    parser.add_argument("--row", type=int, default=None, help="Excel row number to export, e.g. 2.")
    parser.add_argument("--email", default="", help="Find row by email instead of first pending row.")
    parser.add_argument("--out", default=str(ROOT_DIR / "agents" / "runtime" / "current_profile.json"), help="Output JSON file.")
    parser.add_argument("--print-only", action="store_true", help="Print profile JSON without writing.")
    args = parser.parse_args()

    if not args.xlsx:
        raise SystemExit("Missing --xlsx and config.json does not contain xlsx_path.")
    xlsx_path = Path(args.xlsx).expanduser()
    if not xlsx_path.exists():
        raise SystemExit(f"XLSX not found: {xlsx_path}")

    record, row_number, sheet = load_row(xlsx_path, args.sheet, args.row, args.email)
    profile = profile_from_record(record)
    metadata = {"xlsx": str(xlsx_path), "sheet": sheet, "row": row_number}

    payload = {"profile": profile, "source": metadata}
    if args.print_only:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0

    out_path = Path(args.out).expanduser()
    write_profile_json(profile, out_path)
    print(f"Wrote: {out_path}")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
