"""Persist one proxy per phone/account and apply it before opening AirWallet."""
from datetime import datetime, timedelta, timezone
import os
from pathlib import Path

import openpyxl

from src.connections.xlsx_connection import COLUMN_ALIASES, PROXIES_HEADERS, excel_write_lock, replace_excel_file
from src.core.airwallet_otpbase import normalize_phone


def parse_proxy(raw):
    parts = str(raw or '').strip().split(':', 3)
    if (len(parts) != 4 or not parts[1].isascii() or not parts[1].isdigit()
            or not all(parts) or not 1 <= int(parts[1]) <= 65535):
        raise ValueError('Proxy cần dạng host:port:user:password, port từ 1 đến 65535.')
    return {'host': parts[0], 'port': int(parts[1]), 'user': parts[2], 'password': parts[3]}


def _columns(sheet, required=()):
    headers = [COLUMN_ALIASES.get(str(cell.value or '').strip().lower(), str(cell.value or '').strip().lower())
               for cell in sheet[1]]
    for header in required:
        if header not in headers:
            sheet.cell(1, len(headers) + 1, header)
            headers.append(header)
    return {name: index + 1 for index, name in enumerate(headers) if name}


def _text(sheet, row, column):
    return str(sheet.cell(row, column).value or '').strip()


def _save(book, path):
    temporary = path.with_name(f'{path.stem}.{os.getpid()}.proxy.tmp.xlsx')
    try:
        book.save(temporary)
        replace_excel_file(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _pool(sheet):
    columns = _columns(sheet, PROXIES_HEADERS)
    proxies = []
    ids, raw_values = set(), set()
    for row in range(2, sheet.max_row + 1):
        raw = _text(sheet, row, columns['proxy'])
        if not raw:
            continue
        parse_proxy(raw)  # Never include credentials in validation errors.
        identity = _text(sheet, row, columns['proxy_id'])
        if identity and identity in ids:
            raise ValueError('Proxies có proxy_id trùng; cần ID riêng cho mỗi proxy.')
        if raw in raw_values:
            raise ValueError('Proxies có proxy trùng; mỗi dòng cần một proxy riêng.')
        if identity:
            ids.add(identity)
        raw_values.add(raw)
        proxies.append({'proxy_row': row, 'raw': raw, 'proxy_id': identity,
                        'status': _text(sheet, row, columns['status']).upper() or 'ACTIVE',
                        'owner': normalize_phone(_text(sheet, row, columns['assigned_phone_number']))})
    next_id = 1
    for proxy in proxies:
        if not proxy['proxy_id']:
            while f'PX-{next_id:06d}' in ids:
                next_id += 1
            proxy['proxy_id'] = f'PX-{next_id:06d}'
            ids.add(proxy['proxy_id'])
            sheet.cell(proxy['proxy_row'], columns['proxy_id'], proxy['proxy_id'])
    return proxies, columns


def reserve_proxy(xlsx, sheet_name, *, row=None, phone_number=None):
    path = Path(xlsx).expanduser()
    with excel_write_lock(path):
        book = openpyxl.load_workbook(path)
        try:
            if sheet_name not in book.sheetnames or 'Proxies' not in book.sheetnames:
                raise ValueError('Cần sheet dữ liệu và tab Proxies trong Excel; chưa thay đổi proxy trên máy.')
            inputs, pool = book[sheet_name], book['Proxies']
            columns = _columns(inputs, ('proxy_id', 'proxy_used'))
            if 'phone_number' not in columns:
                raise ValueError('Sheet dữ liệu cần cột phone_number để gán proxy theo tài khoản.')
            phone = normalize_phone(phone_number or (_text(inputs, row, columns['phone_number']) if row else ''))
            if not phone:
                raise ValueError('Thiếu số điện thoại để gán proxy.')
            matching = [number for number in range(2, inputs.max_row + 1)
                        if normalize_phone(_text(inputs, number, columns['phone_number'])) == phone]
            if len(matching) != 1 or (row is not None and matching[0] != row):
                raise ValueError('Cần đúng một dòng Excel khớp số điện thoại của tài khoản.')
            row = matching[0]
            proxies, proxy_columns = _pool(pool)
            requested = _text(inputs, row, columns['proxy_id'])
            owned = [proxy for proxy in proxies if proxy['owner'] == phone]
            if len(owned) > 1:
                raise ValueError('Tài khoản đang được gán nhiều proxy; cần kiểm tra tab Proxies.')
            if owned and requested and owned[0]['proxy_id'] != requested:
                raise ValueError('proxy_id của tài khoản không khớp proxy đã gán trước đó; không tự đổi proxy.')
            used_elsewhere = {_text(inputs, number, columns['proxy_id']) for number in range(2, inputs.max_row + 1) if number != row}
            if requested:
                selected = next((proxy for proxy in proxies if proxy['proxy_id'] == requested), None)
                if selected is None:
                    raise ValueError('proxy_id của dòng không có trong tab Proxies.')
            elif owned:
                selected = owned[0]
            else:
                selected = next((proxy for proxy in proxies if not proxy['owner']
                                 and proxy['status'] == 'ACTIVE' and proxy['proxy_id'] not in used_elsewhere), None)
            if selected is None:
                raise ValueError('Hết proxy chưa gán; thêm proxy vào tab Proxies. Không xoay vòng dùng lại.')
            if (selected['owner'] not in ('', phone) or selected['proxy_id'] in used_elsewhere
                    or (selected['status'] == 'ASSIGNED' and not selected['owner'])
                    or selected['status'] not in ('ACTIVE', 'ASSIGNED')):
                raise ValueError('Proxy đã thuộc tài khoản khác hoặc bị tắt; không tự dùng proxy khác.')
            parsed = parse_proxy(selected['raw'])
            inputs.cell(row, columns['proxy_id'], selected['proxy_id'])
            inputs.cell(row, columns['proxy_used'], f"{parsed['host']}:{parsed['port']}")
            pool.cell(selected['proxy_row'], proxy_columns['assigned_phone_number'], phone)
            pool.cell(selected['proxy_row'], proxy_columns['status'], 'ASSIGNED')
            _save(book, path)
            return {**selected, 'input_row': row, 'phone_number': phone}
        finally:
            book.close()


def record_check(xlsx, proxy_id, result):
    path = Path(xlsx).expanduser()
    with excel_write_lock(path):
        book = openpyxl.load_workbook(path)
        try:
            pool = book['Proxies']
            columns = _columns(pool, PROXIES_HEADERS)
            rows = [row for row in range(2, pool.max_row + 1) if _text(pool, row, columns['proxy_id']) == proxy_id]
            if len(rows) != 1:
                raise ValueError('Không xác định duy nhất proxy để ghi kết quả kiểm tra.')
            row = rows[0]
            pool.cell(row, columns['exit_ip'], result.get('ip') or '')
            pool.cell(row, columns['last_checked_at'], datetime.now(timezone(timedelta(hours=7))).strftime('%Y-%m-%d %H:%M:%S +07:00'))
            pool.cell(row, columns['message'], 'Proxy enabled; exit IP verified' if result.get('ok') else str(result.get('error') or 'Proxy check failed'))
            _save(book, path)
        finally:
            book.close()


def prepare_account_proxy(xlsx, sheet_name, device_id, *, row=None, phone_number=None):
    from scripts.proxy_switcher import Phone, apply_proxy
    assignment = reserve_proxy(xlsx, sheet_name, row=row, phone_number=phone_number)
    try:
        result = apply_proxy(Phone(device_id), parse_proxy(assignment['raw']))
    except Exception:
        result = {'ok': False, 'error': 'Không áp dụng được proxy qua ProxySwitcher-NG; giữ nguyên proxy đã gán.'}
    record_check(xlsx, assignment['proxy_id'], result)
    if not result.get('ok'):
        raise RuntimeError(result.get('error') or 'Proxy chưa sẵn sàng; chưa mở AirWallet.')
    return {'proxy_id': assignment['proxy_id'], 'proxy_used': f"{parse_proxy(assignment['raw'])['host']}:{parse_proxy(assignment['raw'])['port']}",
            'ip': result['ip']}


def prepare_profile_proxy(profile, args, settings):
    raw = getattr(args, 'input_xlsx', None) or getattr(args, 'xlsx', None) or settings.get('xlsx_path')
    if not raw:
        raise ValueError('Cần Excel có Inputs và Proxies để giữ một proxy cho mỗi tài khoản.')
    path = Path(raw).expanduser()
    if not path.is_absolute():
        path = Path(__file__).resolve().parents[2] / path
    sheet = getattr(args, 'input_sheet', None) or getattr(args, 'sheet', None) or settings.get('active_sheet', 'Inputs')
    row = getattr(args, 'input_row', None) or getattr(args, 'row', None)
    return prepare_account_proxy(path, sheet, args.device_id, row=row, phone_number=profile['phone_number'])
